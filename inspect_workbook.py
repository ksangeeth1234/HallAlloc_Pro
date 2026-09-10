from openpyxl import load_workbook

path = 'C:/xampp/htdocs/php/Lecture Hall Allocation System/Lecture_Hall_Allocation_for_ongoing_programs.xlsx'
wb = load_workbook(path, data_only=False)
print('SHEETS:', wb.sheetnames)

for ws_name in ['Hall Availability Search', 'Class Schedule Analysis', 'Lecture Hall Availability', 'Dashboard']:
    if ws_name in wb.sheetnames:
        ws = wb[ws_name]
        print('\n---', ws.title, '---')
        print('rows', ws.max_row, 'cols', ws.max_column)
        for row in ws.iter_rows(min_row=1, max_row=min(20, ws.max_row), values_only=True):
            print(row)

print('\n--- DONE ---')
