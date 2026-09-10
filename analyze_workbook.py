from openpyxl import load_workbook

path = 'c:/xampp/htdocs/php/Lecture Hall Allocation System/Lecture_Hall_Allocation_for_ongoing_programs.xlsx'
wb = load_workbook(path, data_only=False)
print('SHEETS:', wb.sheetnames)
for ws in wb.worksheets:
    print(f'\n--- {ws.title} ---')
    print('max_row=', ws.max_row, 'max_col=', ws.max_column)
    for row in ws.iter_rows(min_row=1, max_row=min(12, ws.max_row), values_only=True):
        print(row)

print('\nSheet titles complete.\n')
